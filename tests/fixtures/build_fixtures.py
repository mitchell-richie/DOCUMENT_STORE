"""Generate synthetic test fixtures. Contains no real case material.

Regenerate with:
    uv run python tests/fixtures/build_fixtures.py
"""

from __future__ import annotations

from pathlib import Path

import pymupdf
from docx import Document

FIXTURES_DIR = Path(__file__).resolve().parent


def _render_text_png(text: str) -> bytes:
    """Render text to PNG bytes, used to simulate a scanned page."""
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 72), text, fontsize=18)
    png = page.get_pixmap(dpi=150).tobytes("png")
    doc.close()
    return png


def build_native_pdf(path: Path) -> None:
    doc = pymupdf.open()
    for number in (1, 2):
        page = doc.new_page()
        page.insert_text((72, 72), f"Synthetic native page {number}.", fontsize=14)
        page.insert_text(
            (72, 100),
            "1. The parties agreed to meet on 3 May 2024.",
            fontsize=12,
        )
    doc.save(path)
    doc.close()


def build_scanned_pdf(path: Path) -> None:
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_image(
        page.rect,
        stream=_render_text_png("Synthetic scanned letter. Reference 54321."),
    )
    doc.save(path)
    doc.close()


def build_docx(path: Path) -> None:
    document = Document()
    document.add_heading("Synthetic Heading", level=1)
    document.add_paragraph("1. The first numbered paragraph.")
    document.add_paragraph("2. The second numbered paragraph.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Item"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "Synthetic"
    table.cell(1, 1).text = "42"
    document.save(path)


THREAD_ORIGINAL = """\
From: Alex Example <alex@example.com>
To: Sam Sample <sam@example.org>
Subject: Synthetic meeting request
Date: Mon, 06 May 2024 10:00:00 +0100
Message-ID: <synthetic-1@example.com>
MIME-Version: 1.0
Content-Type: text/plain; charset="utf-8"

Dear Sam,

Please confirm the meeting on 10 May 2024.

Regards,
Alex
"""

THREAD_REPLY = """\
From: Sam Sample <sam@example.org>
To: Alex Example <alex@example.com>
Subject: Re: Synthetic meeting request
Date: Mon, 06 May 2024 14:30:00 +0100
Message-ID: <synthetic-2@example.org>
In-Reply-To: <synthetic-1@example.com>
References: <synthetic-1@example.com>
MIME-Version: 1.0
Content-Type: text/plain; charset="utf-8"

Confirmed for 10 May 2024.

Sam

On Mon, 6 May 2024 at 10:00, Alex Example <alex@example.com> wrote:
> Dear Sam,
>
> Please confirm the meeting on 10 May 2024.
"""


def build_receipt_png(path: Path) -> None:
    path.write_bytes(_render_text_png("Synthetic Receipt  Total 12.50  Date 01/05/2024"))


def build_bank_statement(path: Path) -> None:
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 72), "Synthetic bank statement (not real).", fontsize=14)
    page.insert_text((72, 100), "Opening balance 1000.00  Closing balance 900.00", fontsize=12)
    doc.save(path)
    doc.close()


def main() -> None:
    build_native_pdf(FIXTURES_DIR / "native.pdf")
    build_scanned_pdf(FIXTURES_DIR / "scanned.pdf")
    build_docx(FIXTURES_DIR / "sample.docx")
    (FIXTURES_DIR / "thread_original.eml").write_text(THREAD_ORIGINAL, encoding="utf-8")
    (FIXTURES_DIR / "thread_reply.eml").write_text(THREAD_REPLY, encoding="utf-8")
    build_receipt_png(FIXTURES_DIR / "receipt.png")
    build_bank_statement(FIXTURES_DIR / "bank_statement_synthetic_2024-01.pdf")
    print(f"Fixtures written to {FIXTURES_DIR}")


if __name__ == "__main__":
    main()