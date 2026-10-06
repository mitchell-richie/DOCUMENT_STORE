"""Tests for DOCX extraction (STORY-3.6)."""

from pathlib import Path

import pytest
from docx import Document

from document_store.extract.base import Extractor
from document_store.extract.blocks import Block, Heading, Paragraph, TableRow
from document_store.extract.docx import DocxExtractor, extract_docx
from document_store.extract.errors import ExtractionError
from document_store.extract.pdf import PdfExtractor


def test_fixture_blocks_are_in_document_order(fixtures_dir: Path) -> None:
    blocks = extract_docx(fixtures_dir / "sample.docx").blocks

    assert [type(b) for b in blocks] == [Heading, Paragraph, Paragraph, TableRow, TableRow]
    assert [b.ordinal for b in blocks] == [1, 2, 3, 4, 5]


def test_heading_has_level_and_text(fixtures_dir: Path) -> None:
    heading = extract_docx(fixtures_dir / "sample.docx").blocks[0]

    assert isinstance(heading, Heading)
    assert heading.text == "Synthetic Heading"
    assert heading.level == 1


def test_table_rows_carry_markers_and_cells(fixtures_dir: Path) -> None:
    blocks = extract_docx(fixtures_dir / "sample.docx").blocks
    rows = [b for b in blocks if isinstance(b, TableRow)]

    assert [(r.table_index, r.row_index, r.text) for r in rows] == [
        (1, 1, "Item | Value"),
        (1, 2, "Synthetic | 42"),
    ]


def test_paragraph_has_no_heading_or_table_fields(fixtures_dir: Path) -> None:
    paragraph = extract_docx(fixtures_dir / "sample.docx").blocks[1]

    assert isinstance(paragraph, Paragraph)
    assert not hasattr(paragraph, "level")
    assert not hasattr(paragraph, "table_index")


def test_table_between_paragraphs_keeps_order(tmp_path: Path) -> None:
    target = tmp_path / "mixed.docx"
    document = Document()
    document.add_paragraph("Before the table.")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "A"
    table.cell(0, 1).text = "B"
    document.add_paragraph("After the table.")
    document.save(target)

    blocks = extract_docx(target).blocks
    assert [type(b) for b in blocks] == [Paragraph, TableRow, Paragraph]
    assert blocks[2].text == "After the table."


def test_empty_paragraphs_and_rows_are_skipped(tmp_path: Path) -> None:
    target = tmp_path / "blanks.docx"
    document = Document()
    document.add_paragraph("")
    document.add_paragraph("   ")
    document.add_paragraph("Only content.")
    document.save(target)

    assert [b.text for b in extract_docx(target).blocks] == ["Only content."]


def test_no_page_number_is_recorded(fixtures_dir: Path) -> None:
    block = extract_docx(fixtures_dir / "sample.docx").blocks[0]
    assert not hasattr(block, "page")


def test_unreadable_file_raises(tmp_path: Path) -> None:
    target = tmp_path / "broken.docx"
    target.write_bytes(b"not a docx")

    with pytest.raises(ExtractionError, match="cannot open Word document"):
        extract_docx(target)


def test_extractors_share_the_base_class() -> None:
    assert issubclass(DocxExtractor, Extractor)
    assert issubclass(PdfExtractor, Extractor)


def test_blocks_share_the_base_class() -> None:
    assert issubclass(Heading, Block)
    assert issubclass(Paragraph, Block)
    assert issubclass(TableRow, Block)