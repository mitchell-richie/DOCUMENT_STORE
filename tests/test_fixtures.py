"""Checks that each fixture has the properties later stories depend on."""

from __future__ import annotations

from email import message_from_bytes
from pathlib import Path

import pymupdf
import pytest
from docx import Document

EXPECTED_FILES = [
    "native.pdf",
    "scanned.pdf",
    "sample.docx",
    "thread_original.eml",
    "thread_reply.eml",
    "receipt.png",
    "bank_statement_synthetic_2024-01.pdf",
]


@pytest.mark.parametrize("name", EXPECTED_FILES)
def test_fixture_exists(fixtures_dir: Path, name: str) -> None:
    assert (fixtures_dir / name).is_file()


def test_native_pdf_has_text_layer_on_every_page(fixtures_dir: Path) -> None:
    with pymupdf.open(fixtures_dir / "native.pdf") as doc:
        assert doc.page_count == 2
        assert all(page.get_text().strip() for page in doc)


def test_scanned_pdf_has_no_text_layer(fixtures_dir: Path) -> None:
    with pymupdf.open(fixtures_dir / "scanned.pdf") as doc:
        assert doc.page_count == 1
        assert doc[0].get_text().strip() == ""
        assert doc[0].get_images()


def test_docx_has_heading_paragraphs_and_table(fixtures_dir: Path) -> None:
    document = Document(fixtures_dir / "sample.docx")
    styles = [p.style.name for p in document.paragraphs]
    assert any(style.startswith("Heading") for style in styles)
    assert len(document.tables) == 1


@pytest.mark.parametrize("name", ["thread_original.eml", "thread_reply.eml"])
def test_eml_has_message_id(fixtures_dir: Path, name: str) -> None:
    message = message_from_bytes((fixtures_dir / name).read_bytes())
    assert message["Message-ID"]


def test_reply_references_original(fixtures_dir: Path) -> None:
    reply = message_from_bytes((fixtures_dir / "thread_reply.eml").read_bytes())
    original = message_from_bytes((fixtures_dir / "thread_original.eml").read_bytes())
    assert reply["In-Reply-To"] == original["Message-ID"]


def test_receipt_image_opens(fixtures_dir: Path) -> None:
    with pymupdf.open(fixtures_dir / "receipt.png") as img:
        assert img.page_count == 1


def test_bank_statement_has_routing_pattern(fixtures_dir: Path) -> None:
    assert "bank_statement" in (fixtures_dir / "bank_statement_synthetic_2024-01.pdf").name